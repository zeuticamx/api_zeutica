# Tests de traspasos informativos entre tarjetas (routers/productos.py):
# mismo SKU, origen/destino existentes y distintos, stock suficiente.
# Solo tocan stock_ubicacion + ubicaciones_traspasos; productos queda intacto.
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient

from routers import productos


class FakeCursor:
    def __init__(self, db):
        self.db = db
        self._resultado = []
        self.rowcount = 0

    def execute(self, query, params=None):
        q = " ".join(query.split())
        db = self.db
        db.ejecutados.append((q, params))
        self._resultado, self.rowcount = [], 0
        if q.startswith("CREATE TABLE IF NOT EXISTS ubicaciones_traspasos"):
            return
        if q.startswith("SELECT id, sku, warehouse_id, cantidad FROM stock_ubicacion WHERE id"):
            fila = db.tarjetas.get(params[0])
            self._resultado = [dict(fila)] if fila else []
        elif q.startswith("UPDATE stock_ubicacion SET cantidad"):
            fila = db.tarjetas.get(params[1])
            if fila:
                fila["cantidad"] = params[0]
                db.updates.append((params[1], params[0]))
                self.rowcount = 1
        elif q.startswith("INSERT INTO ubicaciones_traspasos"):
            db.inserts.append(params)
            self.rowcount = 1
        elif q.startswith("SELECT * FROM ubicaciones_traspasos WHERE sku"):
            self._resultado = [dict(f) for f in db.historial if f["sku"] == params[0]]
        elif "FROM stock_ubicacion s LEFT JOIN productos p" in q:
            filas = db.reporte_filas
            if "WHERE s.sku = %s" in q:
                filas = [f for f in filas if f["sku"] == params[0]]
            self._resultado = [dict(f) for f in filas]

    def fetchone(self):
        return self._resultado[0] if self._resultado else None

    def fetchall(self):
        return list(self._resultado)

    def close(self):
        pass


class FakeDB:
    def __init__(self):
        # Mismo SKU en dos tarjetas; tercera de otro SKU para el test de mezcla.
        self.tarjetas = {
            1: {"id": 1, "sku": "BOLECO17", "warehouse_id": "CEDIS-A1", "cantidad": 10},
            2: {"id": 2, "sku": "BOLECO17", "warehouse_id": "CEDIS-B2", "cantidad": 4},
            3: {"id": 3, "sku": "OTRO", "warehouse_id": "CEDIS-A1", "cantidad": 99},
        }
        self.historial = [
            {"id": 1, "sku": "BOLECO17", "origen_warehouse_id": "CEDIS-A1",
             "destino_warehouse_id": "CEDIS-B2", "cantidad": 2, "usuario": "ana"},
        ]
        self.updates = []
        self.inserts = []
        self.ejecutados = []
        self.commits = 0
        self.rollbacks = 0
        self.reporte_filas = [
            {"sku": "ESPFARBLA", "warehouse_id": "CEDIS-A1", "cama": 1, "cantidad": 5, "nombre": "Esponja Farbla"},
            {"sku": "BOLECO17", "warehouse_id": "CEDIS-B2", "cama": 0, "cantidad": 3, "nombre": "Bolsa Eco"},
            {"sku": "GELALC-250", "warehouse_id": "Oficina", "cama": 0, "cantidad": 7, "nombre": "Gel"},
        ]

    def cursor(self, dictionary=False):
        return FakeCursor(self)

    def commit(self):
        self.commits += 1

    def rollback(self):
        self.rollbacks += 1

    def close(self):
        pass


@pytest.fixture
def db(monkeypatch):
    base = FakeDB()
    monkeypatch.setattr(productos, "get_db_connection", lambda: base)
    movimientos = []
    monkeypatch.setattr(productos.mov_reg, "registrar_movimiento", lambda *a, **k: movimientos.append(a))
    base.movimientos = movimientos
    return base


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(productos.router)
    return TestClient(app)


def base_payload(**extra):
    p = {"sku": "BOLECO17", "origen_id": 1, "destino_id": 2, "cantidad": 3, "usuario": "ana"}
    p.update(extra)
    return p


def test_ruta_feliz_resta_origen_suma_destino_y_registra(client, db):
    r = client.post("/ubicaciones/traspaso", json=base_payload())
    assert r.status_code == 200
    body = r.json()
    assert body["origen"] == {"id": 1, "warehouse_id": "CEDIS-A1", "cantidad": 7}
    assert body["destino"] == {"id": 2, "warehouse_id": "CEDIS-B2", "cantidad": 7}
    assert db.updates == [(1, 7), (2, 7)] and len(db.inserts) == 1
    assert db.inserts[0][:6] == ("BOLECO17", 1, "CEDIS-A1", 2, "CEDIS-B2", 3)
    assert db.commits == 1
    assert any("Traspasó 3" in m[1] and m[0] == "ana" for m in db.movimientos)
    # Nada toca productos: ningún UPDATE/SELECT sobre esa tabla.
    assert all("productos" not in q for q, _ in db.ejecutados)


def test_mismo_origen_y_destino_da_422_sin_escribir(client, db):
    r = client.post("/ubicaciones/traspaso", json=base_payload(destino_id=1))
    assert r.status_code == 422
    assert db.updates == [] and db.inserts == [] and db.commits == 0


@pytest.mark.parametrize("payload", [
    base_payload(cantidad=0),
    base_payload(cantidad=-2),
    base_payload(origen_id=0),
    base_payload(sku=""),
    base_payload(usuario=""),
])
def test_payload_invalido_da_422_sin_escribir(client, db, payload):
    assert client.post("/ubicaciones/traspaso", json=payload).status_code == 422
    assert db.updates == [] and db.inserts == [] and db.commits == 0


def test_stock_insuficiente_da_422_sin_escribir(client, db):
    r = client.post("/ubicaciones/traspaso", json=base_payload(cantidad=11))
    assert r.status_code == 422 and "disponible 10" in r.json()["detail"]
    assert db.updates == [] and db.inserts == [] and db.commits == 0


def test_tarjeta_inexistente_da_404(client, db):
    assert client.post("/ubicaciones/traspaso", json=base_payload(origen_id=999)).status_code == 404
    assert client.post("/ubicaciones/traspaso", json=base_payload(destino_id=999)).status_code == 404
    assert db.updates == [] and db.inserts == []


def test_mezclar_skus_da_422(client, db):
    r = client.post("/ubicaciones/traspaso", json=base_payload(destino_id=3))
    assert r.status_code == 422 and "mismo SKU" in r.json()["detail"]
    assert db.updates == [] and db.inserts == []


def test_historial_por_sku_y_vacio_da_lista_vacia(client, db):
    r = client.get("/ubicaciones/traspasos/BOLECO17")
    assert r.status_code == 200 and r.json() == db.historial
    r = client.get("/ubicaciones/traspasos/SIN-NADA")
    assert r.status_code == 200 and r.json() == []


def test_multiplos_no_divergen_del_panel():
    # Espejo de MULTIPLOS_POR_SKU en panel_gerencia/src/ubicaciones-logica.js.
    assert productos.MULTIPLOS_UBICACIONES_POR_SKU == {
        "CUBMANBLA": 20, "CUBBARBLA": 20, "BOLECO17": 20, "BOLECO38": 5, "ESPFARBLA": 1,
    }
    assert productos.MULTIPLO_UBICACIONES_DEFAULT == 10
    assert productos.multiplo_ubicacion("espfarbla", "CEDIS-A1") == 1
    assert productos.multiplo_ubicacion("BOLECO17", "Oficina") == 1
    assert productos.multiplo_ubicacion("X", "CEDIS-E5") == 10


def _libro(r):
    import io
    import openpyxl
    assert r.status_code == 200
    assert "spreadsheet" in r.headers["content-type"]
    return openpyxl.load_workbook(io.BytesIO(r.content)).active


def test_reporte_conteo_xlsx_con_formulas_y_multiplos(client, db):
    ws = _libro(client.get("/ubicaciones/reporte-conteo"))
    assert [c.value for c in ws[3]] == ["SKU", "Producto", "Ubicación", "Cama",
                                        "Cant. sistema (bultos)", "Múltiplo", "Piezas sistema (ref)",
                                        "Conteo 1", "Conteo 2", "Diferencia", "Observaciones"]
    filas = {ws.cell(row=i, column=1).value: i for i in range(4, 7)}
    assert ws.cell(row=filas["ESPFARBLA"], column=6).value == 1
    assert ws.cell(row=filas["ESPFARBLA"], column=7).value == 5
    assert ws.cell(row=filas["BOLECO17"], column=6).value == 20
    assert ws.cell(row=filas["BOLECO17"], column=7).value == 60
    assert ws.cell(row=filas["GELALC-250"], column=6).value == 1  # Oficina siempre x1
    for i in range(4, 7):
        dif = ws.cell(row=i, column=10).value
        assert isinstance(dif, str) and dif.startswith("=IF(") and f"H{i}-E{i}" in dif
        assert ws.cell(row=i, column=8).value in (None, "")
    assert ws.cell(row=7, column=1).value == "TOTALES"
    assert ws.cell(row=7, column=5).value == "=SUM(E4:E6)"
    assert "todos" in client.get("/ubicaciones/reporte-conteo").headers["content-disposition"]


def test_reporte_conteo_filtra_por_sku(client, db):
    r = client.get("/ubicaciones/reporte-conteo", params={"sku": "ESPFARBLA"})
    ws = _libro(r)
    assert ws.cell(row=4, column=1).value == "ESPFARBLA"
    assert ws.cell(row=5, column=1).value == "TOTALES"
    assert "ESPFARBLA" in r.headers["content-disposition"]
