# Tests de la exportación Excel del CRM (crm_exportacion.py): rango del mes, escape de celdas y libro.
import io
from datetime import date, datetime

import openpyxl

import crm_exportacion as ex

RESULTADOS = {"contesto": "Contestó"}


def test_rango_es_del_dia_1_a_hoy():
    assert ex.rango_mes_en_curso(date(2026, 10, 19)) == (date(2026, 10, 1), date(2026, 10, 19))
    assert ex.rango_mes_en_curso(date(2026, 10, 1)) == (date(2026, 10, 1), date(2026, 10, 1))


def test_escapar_celda_neutraliza_formulas_y_control():
    assert ex.escapar_celda("=SUM(A1)") == "'=SUM(A1)"
    assert ex.escapar_celda("+52 55") == "'+52 55"
    assert ex.escapar_celda("@x") == "'@x"
    assert ex.escapar_celda("hola\x00 mundo") == "hola mundo"
    assert ex.escapar_celda(None) == ""
    assert ex.escapar_celda("normal") == "normal"


def test_libro_contiene_filas_etiquetas_y_texto_seguro():
    filas = [{
        "fecha": datetime(2026, 10, 5, 10, 30), "vendedor": "ana", "cliente": "Cliente Ana",
        "tipo": "llamada", "resultado": "contesto", "notas": "=HYPERLINK(\"x\")",
        "proxima_accion": "Llamar", "proxima_fecha": date(2026, 10, 9),
        "seguimiento_cerrado": 0, "etapa": "en_seguimiento",
    }]
    datos = ex.construir_libro(filas, date(2026, 10, 1), date(2026, 10, 19), datetime(2026, 10, 19, 9, 0), RESULTADOS)
    ws = openpyxl.load_workbook(io.BytesIO(datos)).active
    assert [c.value for c in ws[3]] == ex.ENCABEZADOS
    fila = [c.value for c in ws[4]]
    assert fila[1:5] == ["ana", "Cliente Ana", "Llamada", "Contestó"]
    assert fila[5] == "'=HYPERLINK(\"x\")"
    assert fila[8] == "Pendiente" and fila[9] == "En seguimiento"
    assert "1 interacciones" in ws["A2"].value


def test_libro_sin_filas_es_valido():
    datos = ex.construir_libro([], date(2026, 10, 1), date(2026, 10, 1), datetime(2026, 10, 1, 9, 0), RESULTADOS)
    assert openpyxl.load_workbook(io.BytesIO(datos)).active["A3"].value == "Fecha"
