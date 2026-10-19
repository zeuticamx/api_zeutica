# Lógica pura de la exportación Excel de la bitácora del CRM: rango del mes en curso,
# escape de celdas y armado del libro. Sin SQL ni FastAPI (se prueba en tests/test_crm_exportacion.py).
import io
import re
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import crm_metricas

ETIQUETAS_TIPO = {"llamada": "Llamada", "correo": "Correo", "whatsapp": "WhatsApp", "reunion": "Reunión"}

ENCABEZADOS = [
    "Fecha", "Vendedor", "Cliente", "Tipo", "Resultado", "Notas",
    "Próxima acción", "Próxima fecha", "Seguimiento", "Etapa actual",
]
ANCHOS = [18, 14, 30, 12, 16, 50, 30, 14, 14, 18]

_FORMULA = ("=", "+", "-", "@")
# Caracteres de control que openpyxl rechaza (IllegalCharacterError); \t \n \r sí se permiten.
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def rango_mes_en_curso(hoy: date):
    """(primer día del mes, hoy): lo único que se puede exportar de la bitácora."""
    return hoy.replace(day=1), hoy


def escapar_celda(valor):
    """Texto libre a celda segura: sin control chars y sin que Excel lo tome por fórmula."""
    if valor is None:
        return ""
    texto = _CONTROL.sub("", str(valor))
    if texto.startswith(_FORMULA):
        return "'" + texto
    return texto


def _estado_seguimiento(fila) -> str:
    if not fila.get("proxima_fecha"):
        return ""
    return "Hecho" if fila.get("seguimiento_cerrado") else "Pendiente"


def construir_libro(filas, desde: date, hasta: date, generado: datetime, resultados: dict) -> bytes:
    """Un libro de una hoja con la bitácora (ya filtrada y ordenada) del rango dado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bitácora CRM"
    ws["A1"] = f"CRM — Bitácora de interacciones {desde.isoformat()} a {hasta.isoformat()}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Generado: {generado.strftime('%Y-%m-%d %H:%M')} · {len(filas)} interacciones"
    ws["A2"].font = Font(italic=True, size=10)

    fila_h = 3
    gris = PatternFill("solid", fgColor="D9D9D9")
    borde = Border(*(Side(style="thin", color="B0B0B0") for _ in range(4)))
    for col, texto in enumerate(ENCABEZADOS, start=1):
        c = ws.cell(row=fila_h, column=col, value=texto)
        c.font = Font(bold=True)
        c.fill = gris
        c.border = borde
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for i, f in enumerate(filas, start=fila_h + 1):
        valores = [
            f.get("fecha"),
            escapar_celda(f.get("vendedor")),
            escapar_celda(f.get("cliente")),
            ETIQUETAS_TIPO.get(f.get("tipo"), escapar_celda(f.get("tipo"))),
            resultados.get(f.get("resultado"), escapar_celda(f.get("resultado"))),
            escapar_celda(f.get("notas")),
            escapar_celda(f.get("proxima_accion")),
            f.get("proxima_fecha"),
            _estado_seguimiento(f),
            crm_metricas.ETIQUETAS_ETAPA.get(f.get("etapa"), escapar_celda(f.get("etapa"))),
        ]
        for col, valor in enumerate(valores, start=1):
            c = ws.cell(row=i, column=col, value=valor)
            c.border = borde
            if col == 1 and valor:
                c.number_format = "yyyy-mm-dd hh:mm"
            elif col == 8 and valor:
                c.number_format = "yyyy-mm-dd"
            elif col == 6:
                c.alignment = Alignment(wrap_text=True, vertical="top")

    for col, ancho in enumerate(ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = ancho
    ws.freeze_panes = f"A{fila_h + 1}"
    ultima = max(fila_h + len(filas), fila_h + 1)
    ws.auto_filter.ref = f"A{fila_h}:{get_column_letter(len(ENCABEZADOS))}{ultima}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
